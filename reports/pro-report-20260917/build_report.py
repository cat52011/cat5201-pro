from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'cat5201_pro_完整專題報告.docx'
QA = ROOT / 'qa'
QA.mkdir(exist_ok=True)
FONT = 'Microsoft JhengHei'

def diagram(name, labels, subtitle):
    im = Image.new('RGB', (1600, 560), 'white')
    d = ImageDraw.Draw(im)
    ft = ImageFont.truetype('C:/Windows/Fonts/msjh.ttc', 31)
    small = ImageFont.truetype('C:/Windows/Fonts/msjh.ttc', 26)
    if name == 'architecture':
        boxes = [(70,30,1530,135),(70,205,770,355),(840,205,1530,355),(70,420,1530,525)]
        texts = ['桌面 WPF   畫布  節點  設定  決策視窗', 'NodeService 與 Provider\n上下文  模型請求  結果呈現', 'cat5201.Core\n代理  編排  工作區  產出', '外部服務與本機資料   模型 API  Drive  JSON  附件  成果']
        for box, txt in zip(boxes,texts):
            d.rounded_rectangle(box, radius=12, fill='#eef3f7', outline='#607c91',width=3)
            d.multiline_text(((box[0]+box[2])/2,(box[1]+box[3])/2),txt,font=ft,fill='#172c3d',anchor='mm',align='center',spacing=12)
        for x in [420,1180]:
            for y1,y2 in [(137,201),(357,416)]:
                d.line((x,y1,x,y2),fill='#607c91',width=4)
                d.polygon([(x-8,y2-12),(x+8,y2-12),(x,y2)],fill='#607c91')
        d.line((775,280,834,280),fill='#607c91',width=4)
    else:
        coords = [(60+510*c,40+210*r,500+510*c,190+210*r) for r in range(2) for c in range(3)]
        for i,(box,txt) in enumerate(zip(coords,labels)):
            d.rounded_rectangle(box,radius=12,fill='#eef3f7',outline='#607c91',width=3)
            d.multiline_text(((box[0]+box[2])/2,(box[1]+box[3])/2),str(i+1)+'  '+txt,font=ft,fill='#172c3d',anchor='mm',align='center',spacing=12)
            if i%3!=2:
                x,y=box[2]+6,(box[1]+box[3])//2
                d.line((x,y,x+57,y),fill='#607c91',width=4)
                d.polygon([(x+57,y),(x+46,y-8),(x+46,y+8)],fill='#607c91')
        d.text((800,495),subtitle,font=small,fill='#435b6d',anchor='mm')
    path=QA/(name+'.png')
    im.save(path)
    return path

diagrams={
 'architecture':diagram('architecture',[],''),
 'execution':diagram('execution',['輸入與預算檢查','意圖與模型選擇','能力執行\n資料入工作區','代理協作\n依條件觸發','合成與產出\n經使用者確認','結果與執行紀錄'],'依編號執行  失敗與取消交由各階段處理'),
 'personalization':diagram('personalization',['使用者設定','偏好檔與記憶\n分別保存','載入執行狀態','提示與路由\n套用個人化','模型與產出服務','觀察成果\n驗證設定效果'],'設定保存不等於模型必然遵從  需以結果驗證')
}

raw=(ROOT/'report.md').read_text(encoding='utf-8')
raw=raw.translate(str.maketrans({'画':'畫','组':'組','织':'織','执':'執','输':'輸','旧':'舊','实':'實','发':'發','与':'與','机':'機','际':'際','态':'態'}))
(ROOT/'report.md').write_text(raw,encoding='utf-8')
pages=raw.split('---PAGE---')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
sec.footer_distance=Cm(1)
sec.different_first_page_header_footer=True
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name]
    st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),FONT)
    st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),FONT)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.35
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(27)
doc.styles['Title'].paragraph_format.space_after=Pt(22)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 2'].font.size=Pt(12.5)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(7)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(10)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('cat5201 pro  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def add_text(text,style=None):
    p=doc.add_paragraph(style=style)
    if text.startswith('圖 '):
        p.style=doc.styles['Caption'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if text.startswith('https://'):
        text=text.replace('/', '/\u200b').replace('?','?\u200b').replace('-','-\u200b')
    p.add_run(text.strip('`'))
    if text.startswith('dotnet '):
        for r in p.runs:r.font.size=Pt(9)
    return p

def table(lines):
    rows=[[c.strip() for c in x.strip().strip('|').split('|')] for x in lines]
    t=doc.add_table(rows=0, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=([4.4,12] if len(rows[0])==2 else [3.0,6.5,6.9])
    if rows[0][0] in ['編號','優先']:widths=[1.5,5.6,9.3]
    if rows[0][0]=='時段':widths=[4.2,5.3,6.9]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=txt
            tcPr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6ED' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));tcPr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(0)
                if j==0 and len(txt)<15:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(9);run.bold=i==0
        trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

toc=[]
for i,page in enumerate(pages,1):
    for line in page.splitlines():
        if line.startswith('# ') and (line.startswith('# 第') or line.startswith('# 附') or line=='# 參考資料'):
            toc.append((line[2:],i))

for page_no,page in enumerate(pages):
    if page_no:doc.add_page_break()
    lines=page.strip().splitlines();i=0
    if page_no==2:
        doc.add_heading('目錄',level=1)
        for title,pn in toc:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(11)
            tabs=p.paragraph_format.tab_stops;tabs.add_tab_stop(Cm(15.8),WD_ALIGN_PARAGRAPH.RIGHT)
            p.add_run(title+'\t'+str(pn))
        add_text('第六章集中說明個人化設計，第九章區分已取得的測試結果與後續評估。附錄提供展示腳本與程式查閱入口。')
        continue
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            batch=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):batch.append(lines[i].strip());i+=1
            table(batch);continue
        if line.startswith('!DIAGRAM '):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(diagrams[line.split()[1]]),width=Cm(16.1))
            p.paragraph_format.keep_with_next=True
        elif line.startswith('# '):
            if page_no==0:
                p=doc.add_paragraph(style='Title');p.paragraph_format.space_before=Pt(75);p.add_run(line[2:])
            else:doc.add_heading(line[2:],level=1)
        elif line.startswith('## '):doc.add_heading(line[3:],level=2)
        else:
            p=add_text(line)
            if page_no==0:
                p.paragraph_format.space_after=Pt(24)
                if len(line)<35:
                    for r in p.runs:r.font.size=Pt(14)

doc.core_properties.title='結合多模型協作與執行追蹤之節點式 AI 工作台設計與實作'
doc.core_properties.subject='cat5201 pro 1.8.0 專題報告'
doc.core_properties.author=''
doc.core_properties.keywords='節點工作流, 個人化, 多模型協作, 執行追蹤'
doc.save(OUT)
print('DOCX',OUT)
print('PLANNED_PAGES',len(pages),'CHARS',len(raw))

